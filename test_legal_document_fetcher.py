"""
Test harness for legal_document_fetcher.py

Covers pure-Python logic (no network or Selenium required):
  - URL/URN parsing and filename generation
  - HTML content extraction (Shadow DOM and regular DOM paths)
  - Content cleaning and title extraction
  - Word document creation and saving
  - FetcherConfig and FetchResult dataclasses
  - get_summary() statistics
  - CLI entry-point behaviour (subprocess tests)

Integration tests require a live Chrome browser and network access.
Run them explicitly with:
    pytest -m integration
"""

import subprocess
import sys
from pathlib import Path

import pytest
from bs4 import BeautifulSoup

from legal_document_fetcher import (
    FetcherConfig,
    FetchResult,
    HTMLContentExtractor,
    LegalDocumentFetcher,
    WordDocumentBuilder,
)


# ── Fixtures ─────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def urls_from_file():
    """Load URLs from recent_laws_urls.txt, skipping blanks and comments."""
    url_file = Path(__file__).parent / "recent_laws_urls.txt"
    with open(url_file, encoding="utf-8") as f:
        return [
            line.strip()
            for line in f
            if line.strip() and not line.strip().startswith("#")
        ]


@pytest.fixture
def config(tmp_path):
    return FetcherConfig(output_dir=str(tmp_path / "legal_documents"))


@pytest.fixture
def fetcher(config):
    return LegalDocumentFetcher(config)


@pytest.fixture
def extractor():
    return HTMLContentExtractor()


@pytest.fixture
def builder():
    return WordDocumentBuilder()


# ── Sample HTML fixtures ──────────────────────────────────────────────────────
SHADOW_DOM_HTML = """\
<html><body>
  <sf-unstructured-legislation-viewer></sf-unstructured-legislation-viewer>
</body></html>
<!-- SHADOW DOM CONTENT -->
<div class="p-2 w-100">
  <h1>Lei nº 11.437, de 17 de março de 2023</h1>
  <p>Artigo 1º — Este é o texto da lei.</p>
  <p>Artigo 2º — Mais texto aqui para atingir o mínimo de caracteres necessários.</p>
</div>
<!-- END SHADOW DOM -->"""

REGULAR_DOM_HTML = """\
<html><body>
  <sf-legislation-articulation-text>
    <h1>Constituição Federal de 1988</h1>
    <p>Nós, os representantes do povo brasileiro, reunidos em Assembléia Nacional Constituinte.</p>
  </sf-legislation-articulation-text>
</body></html>"""

FALLBACK_HTML = """\
<html><body>
  <div class="texto">
    <h2>Lei nº 8.078/1990</h2>
    <p>Código de Defesa do Consumidor.</p>
  </div>
</body></html>"""


# ── FetcherConfig ─────────────────────────────────────────────────────────────
class TestFetcherConfig:
    def test_default_values(self, tmp_path):
        cfg = FetcherConfig(output_dir=str(tmp_path / "out"))
        assert cfg.retry_attempts == 3
        assert cfg.delay_between_requests == 2.0
        assert cfg.use_selenium is True
        assert cfg.selenium_wait_time == 20

    def test_creates_output_dir(self, tmp_path):
        out = tmp_path / "new_dir"
        assert not out.exists()
        FetcherConfig(output_dir=str(out))
        assert out.exists()

    def test_no_create_output_dir(self, tmp_path):
        out = tmp_path / "never_created"
        FetcherConfig(output_dir=str(out), create_output_dir=False)
        assert not out.exists()

    def test_custom_values(self, tmp_path):
        cfg = FetcherConfig(
            output_dir=str(tmp_path / "out"),
            retry_attempts=5,
            delay_between_requests=0.5,
            selenium_wait_time=10,
        )
        assert cfg.retry_attempts == 5
        assert cfg.delay_between_requests == 0.5
        assert cfg.selenium_wait_time == 10


# ── FetchResult ───────────────────────────────────────────────────────────────
class TestFetchResult:
    def test_str_success(self):
        r = FetchResult(
            url="https://example.com",
            success=True,
            law_number="lei_10741_20031001",
            filename="lei_10741_20031001.docx",
            fetch_time=1.49,
        )
        assert "✓" in str(r)
        assert "lei_10741_20031001" in str(r)

    def test_str_failure(self):
        r = FetchResult(
            url="https://example.com",
            success=False,
            error_message="Connection timeout",
            fetch_time=5.0,
        )
        assert "✗" in str(r)
        assert "Connection timeout" in str(r)

    def test_default_fields(self):
        r = FetchResult(url="https://example.com", success=True)
        assert r.law_number == ""
        assert r.filename == ""
        assert r.error_message is None
        assert r.fetch_time == 0.0


# ── URL / URN Parsing ─────────────────────────────────────────────────────────
# Expected outputs for every URL in recent_laws_urls.txt
EXPECTED_LAW_NUMBERS = [
    (
        "https://normas.leg.br/?urn=urn:lex:br:federal:lei:2023-03-17;11437",
        "lei_11437_20230317",
    ),
    (
        "https://normas.leg.br/?urn=urn:lex:br:federal:lei:2023-04-06;11472",
        "lei_11472_20230406",
    ),
    (
        "https://normas.leg.br/?urn=urn:lex:br:federal:lei:2021-01-13;14119",
        "lei_14119_20210113",
    ),
    (
        "https://normas.leg.br/?urn=urn:lex:br:federal:lei:2021-12-08;14260",
        "lei_14260_20211208",
    ),
    (
        "https://normas.leg.br/?urn=urn:lex:br:federal:lei:2021-12-29;14286",
        "lei_14286_20211229",
    ),
    (
        "https://normas.leg.br/?urn=urn:lex:br:federal:lei:2023-12-12;14754",
        "lei_14754_20231212",
    ),
]


class TestURLParsing:
    @pytest.mark.parametrize("url,expected", EXPECTED_LAW_NUMBERS)
    def test_known_urls_from_recent_laws_file(self, fetcher, url, expected):
        assert fetcher.extract_law_number_from_url(url) == expected

    def test_all_file_urls_parse_to_lei_prefix(self, fetcher, urls_from_file):
        """Every URL in recent_laws_urls.txt produces a 'lei_' identifier."""
        assert urls_from_file, "recent_laws_urls.txt is empty"
        for url in urls_from_file:
            result = fetcher.extract_law_number_from_url(url)
            assert result.startswith("lei_"), f"Unexpected identifier for {url}: {result}"

    def test_constitution_url(self, fetcher):
        # The filename stem now encodes the document type from the URN, so a
        # constituicao URN yields a "constituicao_" prefix (not a hardcoded
        # "lei_"). This keeps different document types from colliding.
        url = "https://normas.leg.br/impressao?urn=urn:lex:br:federal:constituicao:1988-10-05;1988"
        assert fetcher.extract_law_number_from_url(url) == "constituicao_1988_19881005"

    def test_fallback_on_url_without_urn(self, fetcher):
        result = fetcher.extract_law_number_from_url("https://example.com/no-urn")
        assert result.startswith("lei_")

    def test_fallback_on_empty_string(self, fetcher):
        result = fetcher.extract_law_number_from_url("")
        assert result.startswith("lei_")


# ── Filename Generation ───────────────────────────────────────────────────────
class TestFilenameGeneration:
    def test_basic_filename(self, fetcher):
        path = fetcher.generate_filename("lei_11437_20230317", "https://example.com")
        assert path.endswith("lei_11437_20230317.docx")

    def test_file_is_in_output_dir(self, fetcher):
        path = fetcher.generate_filename("lei_11437_20230317", "https://example.com")
        assert Path(path).parent == Path(fetcher.config.output_dir)

    def test_sanitizes_spaces_and_slashes(self, fetcher):
        path = fetcher.generate_filename("lei 123/abc", "https://example.com")
        name = Path(path).name
        assert " " not in name
        assert "/" not in name

    def test_deduplicates_with_counter(self, fetcher):
        first = fetcher.generate_filename("lei_99999_20240101", "https://example.com")
        Path(first).touch()
        second = fetcher.generate_filename("lei_99999_20240101", "https://example.com")
        assert first != second
        assert "_1.docx" in second


# ── HTML Content Extraction ───────────────────────────────────────────────────
class TestHTMLContentExtractor:
    def test_shadow_dom_extraction(self, extractor):
        result = extractor.extract_main_content(SHADOW_DOM_HTML)
        assert result is not None
        assert "11.437" in result.get_text()

    def test_regular_dom_selector(self, extractor):
        result = extractor.extract_main_content(REGULAR_DOM_HTML)
        assert result is not None
        assert "Constituição" in result.get_text()

    def test_fallback_div_texto_selector(self, extractor):
        result = extractor.extract_main_content(FALLBACK_HTML)
        assert result is not None
        assert "8.078" in result.get_text()

    def test_body_fallback_when_no_selector_matches(self, extractor):
        html = "<html><body><p>Algum texto legal aqui presente.</p></body></html>"
        result = extractor.extract_main_content(html)
        assert result is not None

    def test_returns_none_on_parse_error(self, extractor):
        # Pass a non-string to trigger an internal exception
        result = extractor.extract_main_content(None)
        assert result is None

    def test_clean_removes_script_tags(self, extractor):
        soup = BeautifulSoup(
            "<div><script>alert(1)</script><p>Texto</p></div>", "html.parser"
        )
        cleaned = extractor.clean_content(soup.find("div"))
        assert cleaned.find("script") is None
        assert "Texto" in cleaned.get_text()

    def test_clean_removes_style_tags(self, extractor):
        soup = BeautifulSoup(
            "<div><style>.cls{color:red}</style><p>Texto</p></div>", "html.parser"
        )
        cleaned = extractor.clean_content(soup.find("div"))
        assert cleaned.find("style") is None

    def test_clean_removes_html_comments(self, extractor):
        soup = BeautifulSoup(
            "<div><!-- remove me --><p>Manter</p></div>", "html.parser"
        )
        cleaned = extractor.clean_content(soup.find("div"))
        assert "remove me" not in str(cleaned)

    def test_clean_removes_empty_paragraphs(self, extractor):
        soup = BeautifulSoup(
            "<div><p></p><p>Conteúdo</p></div>", "html.parser"
        )
        cleaned = extractor.clean_content(soup.find("div"))
        paragraphs = cleaned.find_all("p")
        assert all(p.get_text(strip=True) for p in paragraphs)

    def test_clean_preserves_br_and_img(self, extractor):
        soup = BeautifulSoup(
            "<div><br/><img src='x.png'/><p>Texto</p></div>", "html.parser"
        )
        cleaned = extractor.clean_content(soup.find("div"))
        assert cleaned.find("br") is not None
        assert cleaned.find("img") is not None

    def test_clean_removes_title_tag(self, extractor):
        # normas.leg.br sometimes leaves an MS-Word-exported <title> at the
        # top of the body. It must be stripped, otherwise it leaks into the
        # output as a garbage paragraph (e.g., "LEI Nº 4" for Lei 4117/1962).
        html = (
            "<div><title>LEI Nº 4</title>"
            "<p>CÂMARA DOS DEPUTADOS</p></div>"
        )
        cleaned = extractor.clean_content(BeautifulSoup(html, "html.parser").find("div"))
        assert cleaned.find("title") is None
        assert "LEI Nº 4" not in cleaned.get_text()
        assert "CÂMARA DOS DEPUTADOS" in cleaned.get_text()

    def test_clean_preserves_empty_table_cells(self, extractor):
        # Empty <td>/<th> carry a grid position — removing them shifts the
        # remaining cells into the wrong column. Regression for the Lei
        # 5.070/1966 ANEXO I rows that have no subitem (middle cell empty).
        html = (
            "<div><table>"
            "<tr><th>A</th><th></th><th>C</th></tr>"
            "<tr><td>service</td><td></td><td>value</td></tr>"
            "</table></div>"
        )
        cleaned = extractor.clean_content(BeautifulSoup(html, "html.parser").find("div"))
        rows = cleaned.find_all("tr")
        assert [len(r.find_all(["td", "th"])) for r in rows] == [3, 3]
        assert rows[1].find_all("td")[1].get_text() == ""

    def test_get_title_from_h1(self, extractor):
        soup = BeautifulSoup(
            "<div><h1>Lei nº 11.437/2023</h1><p>Texto</p></div>", "html.parser"
        )
        assert extractor.get_law_title(soup) == "Lei nº 11.437/2023"

    def test_get_title_from_h2_when_no_h1(self, extractor):
        soup = BeautifulSoup(
            "<div><h2>Lei Complementar nº 123</h2><p>Texto</p></div>", "html.parser"
        )
        assert extractor.get_law_title(soup) == "Lei Complementar nº 123"

    def test_get_title_regex_fallback(self, extractor):
        soup = BeautifulSoup(
            "<div><p>Lei nº 10.741, de 1 de outubro de 2003.</p></div>", "html.parser"
        )
        title = extractor.get_law_title(soup)
        assert "Lei" in title

    def test_get_title_generic_fallback(self, extractor):
        soup = BeautifulSoup("<div><p>Sem título aqui.</p></div>", "html.parser")
        assert extractor.get_law_title(soup) == "Legal Document"

    def test_get_title_prefers_law_heading_over_section_heading(self, extractor):
        # Regression: an MS-Word-exported document may have
        # <h1>INTRODUÇÃO</h1> as the first section heading, followed later
        # by a proper <h2> with the law's identity. The title extractor must
        # skip the section heading and pick the law-looking one.
        soup = BeautifulSoup(
            "<div>"
            "<h1>INTRODUÇÃO</h1>"
            "<h2>LEI Nº 4.117, DE 27 DE AGOSTO DE 1962</h2>"
            "<p>Institui o Código Brasileiro de Telecomunicações.</p>"
            "</div>",
            "html.parser",
        )
        title = extractor.get_law_title(soup)
        assert title == "LEI Nº 4.117, DE 27 DE AGOSTO DE 1962"

    def test_get_title_skips_section_heading_and_uses_body_text(self, extractor):
        # Regression for Lei 4117/1962 (real-world shape): the only headings
        # in the body are section titles like "INTRODUÇÃO", while the law's
        # identity sits in a plain <p>. The extractor must fall through the
        # heading pass and pick up the law phrase from the body instead of
        # promoting "INTRODUÇÃO" as the document title.
        soup = BeautifulSoup(
            "<div>"
            "<h1>INTRODUÇÃO</h1>"
            "<p>CÂMARA DOS DEPUTADOS</p>"
            "<p>LEI Nº 4.117, DE 27 DE AGOSTO DE 1962</p>"
            "</div>",
            "html.parser",
        )
        title = extractor.get_law_title(soup)
        assert title != "INTRODUÇÃO"
        assert "Lei" in title or "LEI" in title
        assert "4.117" in title or "4117" in title

    def test_get_title_constituicao_matches_keyword(self, extractor):
        soup = BeautifulSoup(
            "<div>"
            "<h1>CAPÍTULO I</h1>"
            "<h1>Constituição da República Federativa do Brasil de 1988</h1>"
            "</div>",
            "html.parser",
        )
        title = extractor.get_law_title(soup)
        assert "Constituição" in title


def _build_table(builder, html):
    """Render an HTML <table> via WordDocumentBuilder._add_table and return the docx table."""
    from docx import Document
    soup = BeautifulSoup(html, "html.parser")
    doc = Document()
    builder._add_table(doc, soup.find("table"))
    return doc.tables[0] if doc.tables else None


# ── Word Document Builder ─────────────────────────────────────────────────────
class TestWordDocumentBuilder:
    def test_create_document_contains_paragraph_text(self, builder):
        soup = BeautifulSoup("<div><p>Artigo 1º — Texto da lei.</p></div>", "html.parser")
        doc = builder.create_document(soup.find("div"), "Lei Teste")
        full_text = " ".join(p.text for p in doc.paragraphs)
        assert "Artigo" in full_text

    def test_create_document_adds_title_heading(self, builder):
        soup = BeautifulSoup("<div><p>Conteúdo.</p></div>", "html.parser")
        doc = builder.create_document(soup.find("div"), "Minha Lei")
        heading_texts = [
            p.text for p in doc.paragraphs if p.style.name.startswith("Heading")
        ]
        assert "Minha Lei" in heading_texts

    def test_create_document_skips_generic_title(self, builder):
        soup = BeautifulSoup("<div><p>Conteúdo.</p></div>", "html.parser")
        doc = builder.create_document(soup.find("div"), "Legal Document")
        assert not any(
            p.style.name.startswith("Heading") and "Legal Document" in p.text
            for p in doc.paragraphs
        )

    def test_create_document_skips_title_longer_than_200_chars(self, builder):
        soup = BeautifulSoup("<div><p>Texto.</p></div>", "html.parser")
        long_title = "x" * 201
        doc = builder.create_document(soup.find("div"), long_title)
        assert not any(long_title in p.text for p in doc.paragraphs)

    def test_create_document_skips_title_when_body_has_matching_heading(self, builder):
        # Regression for Decreto-Lei 4657/1942: normas.leg.br documents often
        # include the document's own heading inside the body. Promoting the
        # extracted title as a top-level Heading 1 produces a duplicate.
        title = "Lei de Introdução às normas do Direito Brasileiro"
        soup = BeautifulSoup(
            f"<div><p>CÂMARA DOS DEPUTADOS</p><h1>{title}</h1><p>Art. 1º…</p></div>",
            "html.parser",
        )
        doc = builder.create_document(soup.find("div"), title)
        heading_texts = [
            p.text for p in doc.paragraphs if p.style.name.startswith("Heading")
        ]
        # The body's own <h1> must still render, but only once.
        assert heading_texts.count(title) == 1
        # The first paragraph must be the real first body line, not the title.
        assert doc.paragraphs[0].text == "CÂMARA DOS DEPUTADOS"

    def test_create_document_adds_title_when_body_has_no_matching_heading(self, builder):
        # Older documents carry the law identity only in a <p>, not a heading.
        # In that case the promoted Heading 1 must still be emitted.
        title = "Lei nº 5.070, de 7 de julho de 1966"
        soup = BeautifulSoup(
            f"<div><p>{title}</p><p>Art. 1º…</p></div>",
            "html.parser",
        )
        doc = builder.create_document(soup.find("div"), title)
        heading_texts = [
            p.text for p in doc.paragraphs if p.style.name.startswith("Heading")
        ]
        assert title in heading_texts

    def test_create_document_skips_title_with_whitespace_or_case_differences(self, builder):
        # The comparison must be whitespace/case/Unicode-tolerant so that
        # trivial reformatting between extraction and rendering does not
        # defeat the duplicate-suppression guard.
        extracted = "Lei nº 4.117, DE 27 DE AGOSTO DE 1962"
        in_body = "LEI Nº  4.117, de 27 de agosto de 1962"  # case + extra space
        soup = BeautifulSoup(
            f"<div><p>Intro</p><h2>{in_body}</h2><p>Art. 1º…</p></div>",
            "html.parser",
        )
        doc = builder.create_document(soup.find("div"), extracted)
        heading_texts = [
            p.text for p in doc.paragraphs if p.style.name.startswith("Heading")
        ]
        # Only the body's heading should be present; the extracted title
        # must not be promoted to a second Heading 1. The body heading is
        # rendered with internal whitespace collapsed (double space -> single),
        # so compare against that normalized form.
        in_body_normalized = "LEI Nº 4.117, de 27 de agosto de 1962"
        assert extracted not in heading_texts
        assert in_body_normalized in heading_texts
        # Exactly one heading (no duplicate promotion).
        assert len(heading_texts) == 1

    def test_headings_converted(self, builder):
        html = "<div><h1>Título</h1><h2>Capítulo</h2><p>Texto.</p></div>"
        soup = BeautifulSoup(html, "html.parser")
        doc = builder.create_document(soup.find("div"), "Headings Test")
        heading_texts = [
            p.text for p in doc.paragraphs if p.style.name.startswith("Heading")
        ]
        assert "Título" in heading_texts or "Capítulo" in heading_texts

    def test_wrapper_tag_does_not_flatten_paragraphs(self, builder):
        # Some source pages (e.g. planalto) wrap the entire body in a single
        # <font> holding many <p>. The walker must recurse into such transparent
        # wrappers instead of flattening everything into one giant paragraph.
        html = (
            "<div><font>"
            "<p>Art. 1º Primeiro artigo.</p>"
            "<p>Art. 2º Segundo artigo.</p>"
            "<p>Art. 3º Terceiro artigo.</p>"
            "</font></div>"
        )
        soup = BeautifulSoup(html, "html.parser")
        doc = builder.create_document(soup.find("div"), "Wrapper Test")
        body_paras = [p.text for p in doc.paragraphs if p.text.strip()
                      and not p.style.name.startswith("Heading")]
        assert len(body_paras) == 3
        assert body_paras[0] == "Art. 1º Primeiro artigo."
        assert body_paras[2] == "Art. 3º Terceiro artigo."

    def test_internal_whitespace_collapsed(self, builder):
        # Hard-wrapped source text with embedded newlines/indentation must not
        # leak into the .docx as spurious line breaks; whitespace is collapsed
        # to single spaces so each paragraph is one flowing line.
        html = (
            "<div><p>\n\n\n   A PRESIDENTA\n     DA REPÚBLICA\n\n , no uso\n\t da "
            "atribuição.\n</p></div>"
        )
        soup = BeautifulSoup(html, "html.parser")
        doc = builder.create_document(soup.find("div"), "WS Test")
        para = next(p for p in doc.paragraphs if p.text.strip()
                    and not p.style.name.startswith("Heading"))
        assert "\n" not in para.text
        assert "  " not in para.text  # no double spaces
        assert para.text.strip() == "A PRESIDENTA DA REPÚBLICA , no uso da atribuição."

    def test_bold_formatting_preserved(self, builder):
        html = "<div><p><strong>Texto em negrito</strong> e normal.</p></div>"
        soup = BeautifulSoup(html, "html.parser")
        doc = builder.create_document(soup.find("div"), "Formatting Test")
        bold_runs = [run for p in doc.paragraphs for run in p.runs if run.bold]
        assert any("negrito" in run.text for run in bold_runs)

    def test_save_document_creates_file(self, builder, tmp_path):
        soup = BeautifulSoup("<div><p>Texto de lei.</p></div>", "html.parser")
        doc = builder.create_document(soup.find("div"), "Teste")
        filepath = str(tmp_path / "test_output.docx")
        builder.save_document(doc, filepath)
        assert Path(filepath).exists()
        assert Path(filepath).stat().st_size > 0

    def test_save_document_raises_on_bad_path(self, builder, tmp_path):
        soup = BeautifulSoup("<div><p>Texto.</p></div>", "html.parser")
        doc = builder.create_document(soup.find("div"), "Teste")
        with pytest.raises(IOError):
            builder.save_document(doc, "/nonexistent_dir/output.docx")

    def test_table_plain_three_columns(self, builder):
        html = """
        <table>
          <tr><th>A</th><th>B</th><th>C</th></tr>
          <tr><td>1</td><td>2</td><td>3</td></tr>
        </table>
        """
        table = _build_table(builder, html)
        assert len(table.rows) == 2
        assert len(table.columns) == 3
        assert [c.text for c in table.rows[0].cells] == ["A", "B", "C"]
        assert [c.text for c in table.rows[1].cells] == ["1", "2", "3"]

    def test_table_rowspan_shifts_subsequent_cells(self, builder):
        # Regression for Lei 5.070/1966 ANEXO I: a rowspan in the first column
        # must not push subitems into column 0 or values into column 1.
        html = """
        <table>
          <tr><th>SERVIÇO</th><th>sub</th><th>VALOR</th></tr>
          <tr><td rowspan="3">Telefonia Fixa</td><td>local</td><td>100</td></tr>
          <tr><td>LDN</td><td>200</td></tr>
          <tr><td>LDI</td><td>300</td></tr>
        </table>
        """
        table = _build_table(builder, html)
        col1 = [table.rows[i].cells[1].text for i in range(1, 4)]
        col2 = [table.rows[i].cells[2].text for i in range(1, 4)]
        assert col1 == ["local", "LDN", "LDI"]
        assert col2 == ["100", "200", "300"]
        # Origin cell in column 0 carries the text; python-docx joins the
        # merged cells' paragraphs with newlines, so compare on the stripped
        # first line.
        assert table.rows[1].cells[0].text.splitlines()[0] == "Telefonia Fixa"

    def test_table_colspan_widens_cell(self, builder):
        html = """
        <table>
          <tr><td>A</td><td>B</td><td>C</td></tr>
          <tr><td colspan="2">wide</td><td>X</td></tr>
        </table>
        """
        table = _build_table(builder, html)
        assert len(table.columns) == 3
        assert table.rows[1].cells[0].text.splitlines()[0] == "wide"
        assert table.rows[1].cells[2].text == "X"

    def test_table_rowspan_and_colspan_combined(self, builder):
        html = """
        <table>
          <tr><td rowspan="2" colspan="2">BIG</td><td>c</td></tr>
          <tr><td>f</td></tr>
          <tr><td>g</td><td>h</td><td>i</td></tr>
        </table>
        """
        table = _build_table(builder, html)
        assert len(table.columns) == 3
        assert table.rows[0].cells[0].text.splitlines()[0] == "BIG"
        assert table.rows[0].cells[2].text == "c"
        assert table.rows[1].cells[2].text == "f"
        assert [c.text for c in table.rows[2].cells] == ["g", "h", "i"]

    def test_table_empty_returns_without_error(self, builder):
        from docx import Document
        soup = BeautifulSoup("<table></table>", "html.parser")
        doc = Document()
        builder._add_table(doc, soup.find("table"))
        assert len(doc.tables) == 0

    def test_table_invalid_span_attribute_defaults_to_one(self, builder):
        html = """
        <table>
          <tr><td rowspan="abc">foo</td><td>bar</td></tr>
          <tr><td>baz</td><td>qux</td></tr>
        </table>
        """
        table = _build_table(builder, html)
        assert table.rows[0].cells[0].text == "foo"
        assert table.rows[0].cells[1].text == "bar"
        assert table.rows[1].cells[0].text == "baz"
        assert table.rows[1].cells[1].text == "qux"


# ── get_summary ───────────────────────────────────────────────────────────────
class TestGetSummary:
    def test_empty_results(self, fetcher):
        summary = fetcher.get_summary()
        assert summary["total"] == 0
        assert summary["success"] == 0
        assert summary["failed"] == 0
        assert summary["success_rate"] == 0
        assert summary["avg_fetch_time"] == 0

    def test_all_success(self, fetcher):
        fetcher.results = [
            FetchResult(url="u1", success=True, fetch_time=1.0),
            FetchResult(url="u2", success=True, fetch_time=3.0),
        ]
        summary = fetcher.get_summary()
        assert summary["total"] == 2
        assert summary["success"] == 2
        assert summary["failed"] == 0
        assert summary["success_rate"] == 100.0
        assert summary["avg_fetch_time"] == pytest.approx(2.0)

    def test_mixed_results(self, fetcher):
        fetcher.results = [
            FetchResult(url="u1", success=True, fetch_time=1.0),
            FetchResult(url="u2", success=False, fetch_time=2.0, error_message="err"),
        ]
        summary = fetcher.get_summary()
        assert summary["total"] == 2
        assert summary["success"] == 1
        assert summary["failed"] == 1
        assert summary["success_rate"] == pytest.approx(50.0)
        assert summary["avg_fetch_time"] == pytest.approx(1.5)
        assert "u2" in summary["failed_urls"]

    def test_all_failed(self, fetcher):
        fetcher.results = [
            FetchResult(url="u1", success=False, fetch_time=1.0, error_message="e"),
        ]
        summary = fetcher.get_summary()
        assert summary["success_rate"] == 0.0
        assert "u1" in summary["failed_urls"]


# ── CLI entry-point ───────────────────────────────────────────────────────────
MODULE = str(Path(__file__).parent / "legal_document_fetcher.py")


class TestCLI:
    def test_help_text(self):
        result = subprocess.run(
            [sys.executable, MODULE, "--help"],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 0
        assert "url_file" in result.stdout
        assert "--output-dir" in result.stdout
        assert "--delay" in result.stdout
        assert "--retries" in result.stdout

    def test_missing_positional_arg_exits_nonzero(self):
        result = subprocess.run(
            [sys.executable, MODULE],
            capture_output=True,
            text=True,
        )
        assert result.returncode != 0

    def test_empty_file_exits_1(self, tmp_path):
        url_file = tmp_path / "empty.txt"
        url_file.write_text("\n\n")
        result = subprocess.run(
            [sys.executable, MODULE, str(url_file)],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 1
        assert "No URLs found" in result.stdout

    def test_only_hash_comments_exits_1(self, tmp_path):
        url_file = tmp_path / "comments.txt"
        url_file.write_text("# Comment line\n# Another comment\n")
        result = subprocess.run(
            [sys.executable, MODULE, str(url_file)],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 1
        assert "No URLs found" in result.stdout

    def test_indented_hash_comments_skipped(self, tmp_path):
        """Lines like '  # comment' must be treated as comments, not URLs."""
        url_file = tmp_path / "indented_comments.txt"
        url_file.write_text("# Normal comment\n  # Indented comment\n\n")
        result = subprocess.run(
            [sys.executable, MODULE, str(url_file)],
            capture_output=True,
            text=True,
        )
        assert result.returncode == 1
        assert "No URLs found" in result.stdout

    def test_nonexistent_file_exits_nonzero(self, tmp_path):
        missing = tmp_path / "does_not_exist.txt"
        result = subprocess.run(
            [sys.executable, MODULE, str(missing)],
            capture_output=True,
            text=True,
        )
        assert result.returncode != 0


# ── Integration tests (require Chrome + network) ──────────────────────────────
@pytest.mark.integration
class TestIntegration:
    """Run with: pytest -m integration"""

    def test_fetch_first_url_from_recent_laws(self, fetcher, urls_from_file):
        assert urls_from_file, "recent_laws_urls.txt must not be empty"
        url = urls_from_file[0]
        result = fetcher.process_single_url(url)
        fetcher.cleanup()
        assert result.success, f"Failed to fetch {url}: {result.error_message}"
        out_file = Path(fetcher.config.output_dir) / result.filename
        assert out_file.exists()
        assert out_file.stat().st_size > 0

    def test_process_url_list_from_file(self, fetcher, urls_from_file, tmp_path):
        results = fetcher.process_url_list(urls_from_file, show_progress=False)
        summary = fetcher.get_summary()
        assert summary["total"] == len(urls_from_file)
        assert summary["success"] > 0, "At least one document should succeed"
